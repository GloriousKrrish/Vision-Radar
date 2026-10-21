import sys
import os
import time

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../packages")))

from visionradar.models.database import SessionLocal
from visionradar.services.excel_export import ExcelExportEngine
import openpyxl

def main():
    sys.stdout.reconfigure(line_buffering=True)
    db = SessionLocal()
    try:
        t0 = time.time()
        print("[TEST] Generating workbook for Job 733...", flush=True)
        stream = ExcelExportEngine.generate_workbook(733, db, mode="final")
        dur = time.time() - t0
        data = stream.getvalue()
        print(f"[TEST] Generated workbook in {dur:.2f}s, size: {len(data):,} bytes", flush=True)

        os.makedirs("data/debug", exist_ok=True)
        out_path = "data/debug/test_job_733_direct.xlsx"
        with open(out_path, "wb") as f:
            f.write(data)
        print(f"[TEST] Saved to {out_path}")

        # Verify with openpyxl
        wb = openpyxl.load_workbook(out_path)
        print(f"[TEST] Sheet names: {wb.sheetnames}")
        for sname in wb.sheetnames:
            ws = wb[sname]
            print(f"  - Sheet '{sname}': {ws.max_row} rows, {ws.max_column} cols")

    finally:
        db.close()

if __name__ == "__main__":
    main()
