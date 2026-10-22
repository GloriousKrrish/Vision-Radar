"""
VisionRadar Phase 5.4 Automated Test Suite:
Examination Excel Intelligence Reporting Upgrade & Verification.

Validates:
1. Sheet names, count (10 sheets), and strict deterministic ordering.
2. Enterprise headers, professional styling, and openpyxl readability.
3. Person vs Object segregation (3 human tracks vs 6 physical material objects).
4. Distinct candidate events (4 events) with timecodes, confidence, and uncertainty.
5. Chronological evidence timeline ordering and observation nature classifications.
6. Event review history and human-in-the-loop proctor verdicts (ACCEPTED, REJECTED, PENDING).
7. Empty and partial dataset handling without exceptions or divide-by-zero errors.
8. Preview endpoint API consistency with actual generated workbook.
9. Cryptographic export seal verification and tamper evidence.
10. Strict domain isolation (Examination vs Traffic).
"""

import io
import pytest
import openpyxl
from fastapi.testclient import TestClient

from apps.api.main import app
from visionradar.models.database import SessionLocal
from visionradar.models.entities import (
    ProcessingJob,
    Video,
    ExaminationTrackRecord,
    ExaminationEventRecord,
    ExaminationZoneRecord,
)
from visionradar.services.excel_export import ExcelExportEngine

client = TestClient(app)

EXPECTED_EXAM_SHEETS = [
    "Executive Summary",
    "Monitored Persons & Tracks",
    "Examination Events",
    "Evidence Timeline",
    "Detection Records",
    "Event Review History",
    "Seating Zones",
    "Processing Metadata",
    "Data Dictionary",
    "Export Metadata",
]

EXPECTED_TRAFFIC_SHEETS = [
    "Executive Summary",
    "Processing Metadata",
    "Unique Tracks",
    "Traffic Events",
    "Speed Analytics",
    "Calibration Metadata",
    "Detection Records",
    "Research & Experiments",
    "Data Dictionary",
    "Export Metadata",
]


@pytest.fixture(scope="module")
def db_session():
    db = SessionLocal()
    yield db
    db.close()


def test_examination_workbook_sheet_names_and_order(db_session):
    """Verify that Job 825 generates exactly 10 domain-isolated sheets in the specified order."""
    job = db_session.query(ProcessingJob).filter(ProcessingJob.id == 825).first()
    if not job:
        job = db_session.query(ProcessingJob).filter(ProcessingJob.domain == "examination").first()

    assert job is not None, "Test requires at least one examination job in database"

    stream = ExcelExportEngine.generate_workbook(job.id, db_session, mode="final")
    assert stream is not None
    wb = openpyxl.load_workbook(stream)
    assert wb.sheetnames == EXPECTED_EXAM_SHEETS, f"Expected {EXPECTED_EXAM_SHEETS}, got {wb.sheetnames}"


def test_examination_workbook_headings_and_structure(db_session):
    """Verify key headings in core worksheets."""
    job = db_session.query(ProcessingJob).filter(ProcessingJob.id == 825).first()
    if not job:
        job = db_session.query(ProcessingJob).filter(ProcessingJob.domain == "examination").first()

    stream = ExcelExportEngine.generate_workbook(job.id, db_session, mode="final")
    wb = openpyxl.load_workbook(stream)

    # Monitored Persons & Tracks headings
    ws_tracks = wb["Monitored Persons & Tracks"]
    person_headers = [ws_tracks.cell(1, c).value for c in range(1, 19)]
    assert "Person Track ID" in person_headers
    assert "Entity Classification" in person_headers
    assert "Assigned Role" in person_headers
    assert "Observed Duration (s)" in person_headers
    assert "Active Detections" in person_headers
    assert "Tracking Quality & Limitations" in person_headers

    # Examination Events headings
    ws_events = wb["Examination Events"]
    event_headers = [ws_events.cell(1, c).value for c in range(1, 22)]
    assert "Event Identifier" in event_headers
    assert "Candidate Behavioral Category" in event_headers
    assert "Priority Level" in event_headers
    assert "Evidence Robustness" in event_headers
    assert "Proctor Review Verdict" in event_headers
    assert "Objective Factual Observation" in event_headers

    # Evidence Timeline headings
    ws_timeline = wb["Evidence Timeline"]
    timeline_headers = [ws_timeline.cell(1, c).value for c in range(1, 13)]
    assert "Timeline Seq #" in timeline_headers
    assert "Video Timecode" in timeline_headers
    assert "Observation Nature" in timeline_headers

    # Event Review History headings
    ws_review = wb["Event Review History"]
    review_headers = [ws_review.cell(1, c).value for c in range(1, 12)]
    assert "Audit Entry ID" in review_headers
    assert "Target Event ID" in review_headers
    assert "Proctor Review Verdict" in review_headers
    assert "Disposition Classification" in review_headers


def test_track_and_physical_object_segregation(db_session):
    """Verify that human persons (3) and physical evidence materials (6) are strictly segregated."""
    job = db_session.query(ProcessingJob).filter(ProcessingJob.id == 825).first()
    if not job:
        pytest.skip("Job 825 not present in database")

    stream = ExcelExportEngine.generate_workbook(job.id, db_session, mode="final")
    wb = openpyxl.load_workbook(stream)
    ws = wb["Monitored Persons & Tracks"]

    # Verify rows 2 to 4 are Person #1, #2, #3
    person_ids = [ws.cell(r, 1).value for r in range(2, 5)]
    assert person_ids == ["#1", "#2", "#3"]

    for r in range(2, 5):
        assert ws.cell(r, 2).value == "HUMAN_CANDIDATE"
        assert ws.cell(r, 3).value == "UNKNOWN"
        assert "UNREGISTERED" in str(ws.cell(r, 5).value)

    # Section 2 header at Row 6
    assert "ASSOCIATED PHYSICAL MATERIALS" in str(ws.cell(6, 1).value)

    # Objects table from row 8 to 13
    obj_ids = [ws.cell(r, 1).value for r in range(8, 14)]
    assert obj_ids == ["#1", "#2", "#3", "#4", "#5", "#6"]

    classes = [ws.cell(r, 2).value for r in range(8, 14)]
    assert classes.count("BOOK") == 3
    assert classes.count("CELL PHONE") == 3


def test_distinct_events_and_review_history_reconciliation(db_session):
    """Verify that events sheet contains 4 distinct events and matches Event Review History."""
    job = db_session.query(ProcessingJob).filter(ProcessingJob.id == 825).first()
    if not job:
        pytest.skip("Job 825 not present in database")

    stream = ExcelExportEngine.generate_workbook(job.id, db_session, mode="final")
    wb = openpyxl.load_workbook(stream)

    ws_events = wb["Examination Events"]
    event_ids = [ws_events.cell(r, 1).value for r in range(2, 6)]
    assert len(event_ids) == 4
    assert "evt_prox_c71564f8" in event_ids
    assert "evt_trans_14d46d78" in event_ids
    assert "evt_phone_88049038" in event_ids
    assert "evt_phone_0c7e1881" in event_ids

    # Event Review History audit entries
    ws_review = wb["Event Review History"]
    review_event_ids = [ws_review.cell(r, 2).value for r in range(2, 6)]
    assert review_event_ids == event_ids

    dispositions = {ws_review.cell(r, 2).value: ws_review.cell(r, 7).value for r in range(2, 6)}
    assert dispositions["evt_prox_c71564f8"] == "PENDING"
    assert dispositions["evt_trans_14d46d78"] == "ACCEPTED"
    assert dispositions["evt_phone_88049038"] == "REJECTED"
    assert dispositions["evt_phone_0c7e1881"] == "ACCEPTED"


def test_evidence_timeline_chronological_integrity(db_session):
    """Verify that Evidence Timeline entries are strictly in ascending chronological order."""
    job = db_session.query(ProcessingJob).filter(ProcessingJob.id == 825).first()
    if not job:
        pytest.skip("Job 825 not present in database")

    stream = ExcelExportEngine.generate_workbook(job.id, db_session, mode="final")
    wb = openpyxl.load_workbook(stream)
    ws = wb["Evidence Timeline"]

    timestamps = []
    for r in range(2, ws.max_row + 1):
        val = ws.cell(r, 3).value  # Column 3: Video Timestamp (s)
        if isinstance(val, (int, float)):
            timestamps.append(val)

    assert len(timestamps) > 0
    # Verify non-decreasing sequence
    assert all(timestamps[i] <= timestamps[i + 1] for i in range(len(timestamps) - 1))


def test_empty_examination_dataset_handling(db_session):
    """Verify that an empty examination job (0 tracks, 0 events) exports cleanly."""
    video = db_session.query(Video).first()
    if not video:
        pytest.skip("No video in database")

    empty_exam_job = ProcessingJob(
        video_id=video.id,
        domain="examination",
        status="SUCCEEDED",
        stage="Completed",
        progress_pct=100.0,
        config_json={"module": "examination"}
    )
    db_session.add(empty_exam_job)
    db_session.commit()
    db_session.refresh(empty_exam_job)

    try:
        stream = ExcelExportEngine.generate_workbook(empty_exam_job.id, db_session, mode="final")
        assert stream is not None
        wb = openpyxl.load_workbook(stream)
        assert wb.sheetnames == EXPECTED_EXAM_SHEETS

        # Check Executive Summary zero counts
        ws_exec = wb["Executive Summary"]
        # Monitored persons should be 0
        found_zero = False
        for r in range(12, 28):
            if ws_exec.cell(r, 1).value == "Total Monitored Persons":
                assert ws_exec.cell(r, 2).value == 0
                found_zero = True
        assert found_zero
    finally:
        db_session.delete(empty_exam_job)
        db_session.commit()


def test_partial_examination_dataset_handling(db_session):
    """Verify handling of partial records (track with no events, missing crops)."""
    video = db_session.query(Video).first()
    if not video:
        pytest.skip("No video in database")

    partial_job = ProcessingJob(
        video_id=video.id,
        domain="examination",
        status="SUCCEEDED",
        stage="Completed",
        progress_pct=100.0,
        config_json={"module": "examination"}
    )
    db_session.add(partial_job)
    db_session.commit()
    db_session.refresh(partial_job)

    track = ExaminationTrackRecord(
        job_id=partial_job.id,
        track_id=101,
        entity_type="person",
        assigned_role="UNKNOWN",
        role_source="DEFAULT",
        first_frame=0,
        last_frame=15,
        trajectory_json=[{"frame": 0, "time": 0.0, "bbox": [10, 10, 50, 50]}]
    )
    db_session.add(track)
    db_session.commit()

    try:
        stream = ExcelExportEngine.generate_workbook(partial_job.id, db_session, mode="final")
        assert stream is not None
        wb = openpyxl.load_workbook(stream)
        assert wb.sheetnames == EXPECTED_EXAM_SHEETS

        ws_tracks = wb["Monitored Persons & Tracks"]
        assert ws_tracks.cell(2, 1).value == "#101"
        assert ws_tracks.cell(2, 2).value == "HUMAN_CANDIDATE"
        assert ws_tracks.cell(2, 11).value == "UNASSIGNED"  # Seating Zone
        assert ws_tracks.cell(2, 12).value == 0  # Associated Candidate Events count
    finally:
        db_session.delete(track)
        db_session.delete(partial_job)
        db_session.commit()


def test_export_preview_endpoint_accuracy(db_session):
    """Verify GET /api/v1/jobs/{job_id}/export/excel/preview accurately segregates counts."""
    job = db_session.query(ProcessingJob).filter(ProcessingJob.id == 825).first()
    if not job:
        pytest.skip("Job 825 not present in database")

    resp = client.get("/api/v1/jobs/825/export/excel/preview")
    assert resp.status_code == 200
    data = resp.json()

    assert data["job_id"] == 825
    assert data["domain"] == "examination"
    assert data["status"] == "SUCCEEDED"
    assert data["counts"]["tracks"] == 3  # Strictly human persons
    assert data["counts"]["person_tracks"] == 3
    assert data["counts"]["object_tracks"] == 6
    assert data["counts"]["events"] == 4
    assert data["counts"]["reviewed_events"] == 3
    assert data["counts"]["pending_events"] == 1
    assert data["planned_sheets"] == EXPECTED_EXAM_SHEETS


def test_cryptographic_export_seal_verification(db_session):
    """Verify presence of SHA-256 seal prefixed with VR-EXAM-SEAL-."""
    job = db_session.query(ProcessingJob).filter(ProcessingJob.id == 825).first()
    if not job:
        job = db_session.query(ProcessingJob).filter(ProcessingJob.domain == "examination").first()

    stream = ExcelExportEngine.generate_workbook(job.id, db_session, mode="final")
    wb = openpyxl.load_workbook(stream)
    ws_meta = wb["Export Metadata"]

    seal_found = False
    for r in range(2, ws_meta.max_row + 1):
        if ws_meta.cell(r, 1).value == "Cryptographic Export Seal":
            val = str(ws_meta.cell(r, 2).value)
            assert val.startswith("VR-EXAM-SEAL-"), f"Unexpected seal format: {val}"
            assert len(val) >= 40
            seal_found = True

    assert seal_found, "Cryptographic Export Seal not found in Export Metadata"


def test_traffic_export_unaltered_and_domain_isolated(db_session):
    """Verify that Traffic exports remain completely untouched and domain-segregated."""
    traffic_job = db_session.query(ProcessingJob).filter(
        ProcessingJob.domain == "traffic",
        ProcessingJob.status == "SUCCEEDED"
    ).first()

    if not traffic_job:
        pytest.skip("No traffic job found to verify")

    stream = ExcelExportEngine.generate_workbook(traffic_job.id, db_session, mode="final")
    wb = openpyxl.load_workbook(stream)

    # Must match expected traffic sheets
    assert wb.sheetnames == EXPECTED_TRAFFIC_SHEETS

    # Must NOT contain examination sheets
    assert "Monitored Persons & Tracks" not in wb.sheetnames
    assert "Examination Events" not in wb.sheetnames
    assert "Evidence Timeline" not in wb.sheetnames
    assert "Event Review History" not in wb.sheetnames
    assert "Seating Zones" not in wb.sheetnames
