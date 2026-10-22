"""
Phase 5.8 Regression Test Suite: Strict Traffic-Examination Isolation, Job Integrity, and Lifecycle Recovery.
Validates:
1. Traffic upload and dispatch
2. Examination upload and dispatch
3. Persisted domain integrity
4. Explicit domain selection vs scene recommendation
5. Cross-domain track isolation
6. Cross-domain event isolation
7. Job-specific telemetry & Common Envelope
8. Job-specific report generation guards
9. Job-specific Excel export isolation
10. Cancelled & superseded worker behavior
"""

import pytest
import io
import openpyxl
from fastapi.testclient import TestClient

from apps.api.main import app
from visionradar.models.database import SessionLocal
from visionradar.models.entities import (
    ProcessingJob, Video, Calibration, Track, SpeedMeasurement, Violation,
    ExaminationTrackRecord, ExaminationEventRecord, ExaminationZoneRecord
)
from visionradar.worker.job_worker import is_job_cancelled, _CANCELLED_JOBS

client = TestClient(app)


@pytest.fixture(scope="module")
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture(scope="module")
def setup_fixture_jobs(db):
    """Ensure at least one valid Traffic job and one valid Examination job exist."""
    # Find or create a video
    video = db.query(Video).first()
    if not video:
        video = Video(
            project_id=1,
            filename="highway_test.mp4",
            storage_path="data/raw_videos/Traffic1.mp4",
            sha256_hash="test_hash_fixture",
            duration_sec=10.0,
            fps=30.0,
            width=1920,
            height=1080
        )
        db.add(video)
        db.commit()
        db.refresh(video)

    # 1. Traffic Job
    traffic_job = db.query(ProcessingJob).filter(
        ProcessingJob.domain == "traffic",
        ProcessingJob.status == "SUCCEEDED"
    ).first()
    if not traffic_job:
        traffic_job = ProcessingJob(
            video_id=video.id,
            domain="traffic",
            status="SUCCEEDED",
            stage="Done",
            progress_pct=100.0,
            config_json={"domain": "traffic", "module": "traffic", "speed_limit_kmh": 80.0}
        )
        db.add(traffic_job)
        db.commit()
        db.refresh(traffic_job)

        # Add mock traffic tracks and speed measurements
        trk = Track(
            job_id=traffic_job.id,
            track_id=1,
            vehicle_class="car",
            confidence=0.92,
            first_frame=1,
            last_frame=30,
            trajectory_json=[{"frame_index": 1, "bbox": [100, 100, 200, 200]}]
        )
        db.add(trk)
        db.commit()
        db.refresh(trk)

        sm = SpeedMeasurement(
            track_id=trk.id,
            frame_index=1,
            timestamp=0.033,
            instantaneous_kmh=85.0,
            smoothed_kmh=85.0,
            uncertainty_kmh=2.0
        )
        viol = Violation(
            job_id=traffic_job.id,
            track_id=trk.id,
            vehicle_class="car",
            frame_index=1,
            timestamp=0.033,
            estimated_speed_kmh=85.0,
            speed_limit_kmh=80.0,
            uncertainty_kmh=2.0,
            review_status="ACCEPTED"
        )
        db.add_all([sm, viol])
        db.commit()

    # 2. Examination Job
    exam_job = db.query(ProcessingJob).filter(
        ProcessingJob.domain == "examination",
        ProcessingJob.status == "SUCCEEDED"
    ).first()
    if not exam_job:
        exam_job = ProcessingJob(
            video_id=video.id,
            domain="examination",
            status="SUCCEEDED",
            stage="Done",
            progress_pct=100.0,
            config_json={"domain": "examination", "module": "examination"}
        )
        db.add(exam_job)
        db.commit()
        db.refresh(exam_job)

        # Add mock person track and exam candidate event
        e_trk = ExaminationTrackRecord(
            job_id=exam_job.id,
            track_id=1,
            entity_type="person",
            first_frame=1,
            last_frame=60,
            assigned_zone="ZONE_DESK_01",
            assigned_role="STUDENT",
            role_source="DEFAULT",
            trajectory_json=[{"frame_index": 1, "bbox": [150, 150, 250, 350]}]
        )
        e_evt = ExaminationEventRecord(
            event_id="EVT_TEST_ISOLATION_01",
            job_id=exam_job.id,
            video_id=video.id,
            event_type="UNAUTHORIZED_OBJECT_INTERACTION",
            person_track_id=1,
            start_frame=1,
            end_frame=30,
            start_time=0.033,
            end_time=1.0,
            confidence=0.88,
            uncertainty=0.12,
            review_status="ACCEPTED",
            supporting_objects_json={"priority": "RED", "evidence_strength": "HIGH", "primary_finding": True}
        )
        db.add_all([e_trk, e_evt])
        db.commit()

    return {"traffic_job_id": traffic_job.id, "exam_job_id": exam_job.id, "video_id": video.id}


# ===========================================================================
# 1. Persisted Domain Integrity & Validation
# ===========================================================================
def test_missing_domain_is_rejected_with_http_400(setup_fixture_jobs):
    """Missing domain must never silently default to Traffic or Examination."""
    vid_id = setup_fixture_jobs["video_id"]
    res = client.post(f"/api/v1/videos/{vid_id}/jobs", json={"calibration_id": None})
    assert res.status_code == 400
    assert "Invalid or missing domain" in res.json()["detail"]


def test_invalid_domain_is_rejected_with_http_400(setup_fixture_jobs):
    """Unrecognized domain must be rejected immediately."""
    vid_id = setup_fixture_jobs["video_id"]
    res = client.post(f"/api/v1/videos/{vid_id}/jobs", json={"domain": "smart_city_drone"})
    assert res.status_code == 400
    assert "Explicit domain ('traffic' or 'examination') is strictly required" in res.json()["detail"]


def test_traffic_job_dispatch_persists_traffic_domain(setup_fixture_jobs, db):
    """Traffic dispatch creates a job with domain='traffic' and assigns calibration if available."""
    vid_id = setup_fixture_jobs["video_id"]
    res = client.post(f"/api/v1/videos/{vid_id}/jobs", json={
        "domain": "traffic",
        "config_json": {"speed_limit_kmh": 90.0}
    })
    assert res.status_code == 200
    data = res.json()
    job_id = data["id"]
    assert data["domain"] == "traffic"

    # Verify persisted DB entity
    job = db.query(ProcessingJob).filter(ProcessingJob.id == job_id).first()
    assert job is not None
    assert job.domain == "traffic"
    assert job.config_json["domain"] == "traffic"


def test_examination_job_dispatch_persists_examination_and_strips_calibration(setup_fixture_jobs, db):
    """Examination dispatch persists domain='examination' and never binds traffic calibration."""
    vid_id = setup_fixture_jobs["video_id"]
    res = client.post(f"/api/v1/videos/{vid_id}/jobs", json={
        "domain": "examination",
        "calibration_id": 12345  # Malicious/stale calibration injection
    })
    assert res.status_code == 200
    data = res.json()
    job_id = data["id"]
    assert data["domain"] == "examination"
    assert data["calibration_id"] is None

    # Verify persisted DB entity
    job = db.query(ProcessingJob).filter(ProcessingJob.id == job_id).first()
    assert job is not None
    assert job.domain == "examination"
    assert job.calibration_id is None


# ===========================================================================
# 2. Cross-Domain Track & Telemetry Isolation
# ===========================================================================
def test_traffic_tracks_endpoint_rejects_examination_job(setup_fixture_jobs):
    """Querying /tracks for an Examination job must return HTTP 400."""
    exam_id = setup_fixture_jobs["exam_job_id"]
    res = client.get(f"/api/v1/jobs/{exam_id}/tracks")
    assert res.status_code == 400
    assert "Vehicle tracks are only available for traffic jobs" in res.json()["detail"]


def test_examination_results_endpoint_rejects_traffic_job(setup_fixture_jobs):
    """Querying /examination/results for a Traffic job must return HTTP 400."""
    traffic_id = setup_fixture_jobs["traffic_job_id"]
    res = client.get(f"/api/v1/jobs/{traffic_id}/examination/results")
    assert res.status_code == 400
    assert "Examination results are only available for examination jobs" in res.json()["detail"]


def test_vehicle_speeds_endpoint_rejects_examination_job(setup_fixture_jobs):
    """Querying /vehicle-speeds for an Examination job must return HTTP 400."""
    exam_id = setup_fixture_jobs["exam_job_id"]
    res = client.get(f"/api/v1/jobs/{exam_id}/vehicle-speeds")
    assert res.status_code == 400
    assert "Vehicle speeds are only available for traffic jobs" in res.json()["detail"]


def test_accident_candidates_endpoint_rejects_examination_job(setup_fixture_jobs):
    """Querying /accident-candidates for an Examination job must return HTTP 400."""
    exam_id = setup_fixture_jobs["exam_job_id"]
    res = client.get(f"/api/v1/jobs/{exam_id}/accident-candidates")
    assert res.status_code == 400
    assert "Accident candidates are only available for traffic jobs" in res.json()["detail"]


def test_speed_limit_config_rejects_examination_job(setup_fixture_jobs):
    """Querying or configuring speed limits for an Examination job must return HTTP 400."""
    exam_id = setup_fixture_jobs["exam_job_id"]
    # GET
    res_get = client.get(f"/api/v1/jobs/{exam_id}/speed-limit")
    assert res_get.status_code == 400
    assert "Speed limit configuration is only available for traffic jobs" in res_get.json()["detail"]

    # POST
    res_post = client.post(f"/api/v1/jobs/{exam_id}/speed-limit", json={
        "speed_limit_kmh": 100.0,
        "warning_tolerance_kmh": 5.0
    })
    assert res_post.status_code == 400
    assert "Speed limit configuration is only available for traffic jobs" in res_post.json()["detail"]


# ===========================================================================
# 3. Cross-Domain Event & Violation Isolation
# ===========================================================================
def test_violations_endpoint_rejects_examination_job(setup_fixture_jobs):
    """Querying /violations with an examination job_id must return HTTP 400."""
    exam_id = setup_fixture_jobs["exam_job_id"]
    res = client.get(f"/api/v1/violations?job_id={exam_id}")
    assert res.status_code == 400
    assert "Traffic violations are only available for traffic jobs" in res.json()["detail"]


def test_examination_events_endpoint_rejects_traffic_job(setup_fixture_jobs):
    """Querying /examination/events with a traffic job_id must return HTTP 400."""
    traffic_id = setup_fixture_jobs["traffic_job_id"]
    res = client.get(f"/api/v1/jobs/{traffic_id}/examination/events")
    assert res.status_code == 400
    assert "Examination events are only available for examination jobs" in res.json()["detail"]


# ===========================================================================
# 4. Report Generation & Audit Isolation
# ===========================================================================
def test_traffic_pdf_report_rejects_examination_job(setup_fixture_jobs):
    """Attempting to generate a Traffic PDF report for an Examination job must return HTTP 400."""
    exam_id = setup_fixture_jobs["exam_job_id"]
    res = client.post(f"/api/v1/jobs/{exam_id}/reports/pdf")
    assert res.status_code == 400
    assert "Traffic PDF report is only available for traffic jobs" in res.json()["detail"]


def test_examination_audit_html_rejects_traffic_job(setup_fixture_jobs):
    """Attempting to view Examination Forensic Audit HTML for a Traffic job must return HTTP 400."""
    traffic_id = setup_fixture_jobs["traffic_job_id"]
    res = client.get(f"/api/v1/jobs/{traffic_id}/examination/audit-report/html")
    assert res.status_code == 400
    assert "Forensic audit HTML is only available for examination jobs" in res.json()["detail"]


def test_examination_audit_report_json_rejects_traffic_job(setup_fixture_jobs):
    """Attempting to get Examination Forensic Audit JSON for a Traffic job must return HTTP 400."""
    traffic_id = setup_fixture_jobs["traffic_job_id"]
    res = client.get(f"/api/v1/jobs/{traffic_id}/examination/audit-report")
    assert res.status_code == 400
    assert "Forensic audit report is only available for examination jobs" in res.json()["detail"]


def test_examination_audit_html_succeeds_for_examination_job(setup_fixture_jobs):
    """Examination Forensic Audit HTML renders valid HTML for Examination job."""
    exam_id = setup_fixture_jobs["exam_job_id"]
    res = client.get(f"/api/v1/jobs/{exam_id}/examination/audit-report/html")
    assert res.status_code == 200
    assert "text/html" in res.headers["content-type"]
    assert "VisionRadar Forensic Audit Report" in res.text
    assert f"Job #{exam_id}" in res.text


# ===========================================================================
# 5. Common Envelope Authoritative Serialization
# ===========================================================================
def test_common_results_envelope_uses_persisted_domain(setup_fixture_jobs):
    """Verify get_job_common_results uses authoritative persisted job.domain."""
    traffic_id = setup_fixture_jobs["traffic_job_id"]
    res_t = client.get(f"/api/v1/jobs/{traffic_id}/results")
    assert res_t.status_code == 200
    env_t = res_t.json()
    assert env_t["analysis_module"] == "traffic"

    exam_id = setup_fixture_jobs["exam_job_id"]
    res_e = client.get(f"/api/v1/jobs/{exam_id}/results")
    assert res_e.status_code == 200
    env_e = res_e.json()
    assert env_e["analysis_module"] == "examination"
    assert "tracks" in env_e["payload"]
    assert "candidate_events" in env_e["payload"]


# ===========================================================================
# 6. Job-Specific Excel Export Isolation
# ===========================================================================
def test_excel_export_preview_isolates_domains(setup_fixture_jobs):
    """Preview metadata must strictly reflect the job's domain."""
    traffic_id = setup_fixture_jobs["traffic_job_id"]
    res_t = client.get(f"/api/v1/jobs/{traffic_id}/export/excel/preview")
    assert res_t.status_code == 200
    data_t = res_t.json()
    assert data_t["domain"] == "traffic"
    assert "Speed Analytics" in data_t["planned_sheets"]

    exam_id = setup_fixture_jobs["exam_job_id"]
    res_e = client.get(f"/api/v1/jobs/{exam_id}/export/excel/preview")
    assert res_e.status_code == 200
    data_e = res_e.json()
    assert data_e["domain"] == "examination"
    assert "Examination Events" in data_e["planned_sheets"]
    assert "Monitored Persons & Tracks" in data_e["planned_sheets"]


def test_excel_export_generates_domain_specific_workbooks(setup_fixture_jobs):
    """Exported Excel workbooks must contain domain-exclusive sheets and correct response headers."""
    # 1. Traffic Export
    traffic_id = setup_fixture_jobs["traffic_job_id"]
    res_t = client.get(f"/api/v1/jobs/{traffic_id}/export/excel?mode=final")
    assert res_t.status_code == 200
    assert res_t.headers["X-VisionRadar-Domain"] == "traffic"
    wb_t = openpyxl.load_workbook(io.BytesIO(res_t.content))
    assert "Unique Tracks" in wb_t.sheetnames
    assert "Traffic Events" in wb_t.sheetnames
    assert "Examination Events" not in wb_t.sheetnames

    # 2. Examination Export
    exam_id = setup_fixture_jobs["exam_job_id"]
    res_e = client.get(f"/api/v1/jobs/{exam_id}/export/excel?mode=final")
    assert res_e.status_code == 200
    assert res_e.headers["X-VisionRadar-Domain"] == "examination"
    wb_e = openpyxl.load_workbook(io.BytesIO(res_e.content))
    assert "Examination Events" in wb_e.sheetnames
    assert "Monitored Persons & Tracks" in wb_e.sheetnames
    assert "Traffic Events" not in wb_e.sheetnames
    assert "Speed Analytics" not in wb_e.sheetnames


# ===========================================================================
# 7. Worker Cancellation & Superseded Job State
# ===========================================================================
def test_cancel_processing_job_updates_status_and_registry(setup_fixture_jobs, db):
    """Calling cancel must transition job status to CANCELLED and mark in cancellation registry."""
    vid_id = setup_fixture_jobs["video_id"]
    # Create an active job
    test_job = ProcessingJob(
        video_id=vid_id,
        domain="traffic",
        status="RUNNING",
        stage="Detection",
        progress_pct=25.0,
        config_json={"domain": "traffic"}
    )
    db.add(test_job)
    db.commit()
    db.refresh(test_job)

    # Cancel via API
    res = client.post(f"/api/v1/jobs/{test_job.id}/cancel")
    assert res.status_code == 200
    assert res.json()["status"] == "CANCELLED"

    # Verify DB state
    db.refresh(test_job)
    assert test_job.status == "CANCELLED"
    assert test_job.stage == "Cancelled"
    assert is_job_cancelled(test_job.id) is True


def test_superseded_worker_cancelled_when_new_job_dispatched(setup_fixture_jobs, db):
    """When a new job is dispatched on a video, prior active jobs on that video are cancelled."""
    vid_id = setup_fixture_jobs["video_id"]

    # Create prior running job
    prior_job = ProcessingJob(
        video_id=vid_id,
        domain="traffic",
        status="RUNNING",
        stage="Detection",
        progress_pct=40.0,
        config_json={"domain": "traffic"}
    )
    db.add(prior_job)
    db.commit()
    db.refresh(prior_job)

    # Dispatch new job
    res = client.post(f"/api/v1/videos/{vid_id}/jobs", json={"domain": "traffic"})
    assert res.status_code == 200
    new_job_id = res.json()["id"]

    # Verify prior job is now CANCELLED and marked superseded
    db.refresh(prior_job)
    assert prior_job.status == "CANCELLED"
    assert f"Superseded by new job #{new_job_id}" in prior_job.error_message
    assert is_job_cancelled(prior_job.id) is True
